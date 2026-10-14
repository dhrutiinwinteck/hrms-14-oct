from __future__ import annotations

import logging
from pathlib import Path
from typing import Final
from urllib.parse import unquote, urlparse

from docx import Document
from pypdf import PdfReader

from app.core.config import settings

logger = logging.getLogger(__name__)


class PolicyDocumentService:
    """
    Handles policy document resolution and text extraction.

    Supported:
        - PDF
        - DOCX

    Translation is NOT performed here.
    Database is NOT modified here.
    """

    _SUPPORTED_FILE_SCHEME: Final[str] = "file"

    _SUPPORTED_EXTENSIONS: Final[set[str]] = {
        ".pdf",
        ".docx",
    }

    # ============================================================
    # PUBLIC EXTRACTION METHOD
    # ============================================================

    def extract_text(
        self,
        document_url: str,
    ) -> str:
        """
        Extract text from PDF or DOCX.

        Supported database values:

            /uploads/policies/file.pdf

            uploads/policies/file.pdf

            file.pdf

            http://127.0.0.1:8001/uploads/policies/file.pdf

            file:///C:/folder/file.pdf
        """

        document_path = self._resolve_document_path(
            document_url
        )

        self._validate_file_exists(
            document_path
        )

        extension = (
            document_path.suffix.lower()
        )

        if extension == ".pdf":
            return self._extract_pdf(
                document_path
            )

        if extension == ".docx":
            return self._extract_docx(
                document_path
            )

        raise ValueError(
            "Unsupported policy document type. "
            "Only PDF and DOCX files are supported."
        )

    # ============================================================
    # PDF BACKWARD COMPATIBILITY
    # ============================================================

    def extract_text_from_pdf(
        self,
        document_url: str,
    ) -> str:
        """
        Backward-compatible PDF-only method.
        """

        pdf_path = self._resolve_pdf_path(
            document_url
        )

        return self._extract_pdf(
            pdf_path
        )

    # ============================================================
    # DOCX EXTRACTION
    # ============================================================

    def extract_text_from_docx(
        self,
        document_url: str,
    ) -> str:
        """
        Extract readable text from DOCX.
        """

        document_path = (
            self._resolve_document_path(
                document_url
            )
        )

        if (
            document_path.suffix.lower()
            != ".docx"
        ):
            raise ValueError(
                "The policy document must be a DOCX file."
            )

        return self._extract_docx(
            document_path
        )

    # ============================================================
    # PDF EXTRACTION
    # ============================================================

    def _extract_pdf(
        self,
        pdf_path: Path,
    ) -> str:
        """
        Extract selectable text from PDF.
        """

        self._validate_file_exists(
            pdf_path
        )

        if (
            pdf_path.suffix.lower()
            != ".pdf"
        ):
            raise ValueError(
                "The policy document must be a PDF file."
            )

        try:
            reader = PdfReader(
                str(pdf_path)
            )

        except Exception as exc:
            logger.exception(
                "Unable to open policy PDF: %s",
                pdf_path,
            )

            raise ValueError(
                "The policy PDF could not be opened."
            ) from exc

        # --------------------------------------------------------
        # ENCRYPTED PDF
        # --------------------------------------------------------

        if reader.is_encrypted:

            try:
                decrypt_result = (
                    reader.decrypt("")
                )

            except Exception as exc:

                raise ValueError(
                    "The policy PDF is encrypted and "
                    "cannot be read automatically."
                ) from exc

            if decrypt_result == 0:

                raise ValueError(
                    "The policy PDF is encrypted and "
                    "cannot be read automatically."
                )

        # --------------------------------------------------------
        # EXTRACT PAGES
        # --------------------------------------------------------

        pages: list[str] = []

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):

            try:

                page_text = (
                    page.extract_text()
                    or ""
                )

            except Exception as exc:

                logger.warning(
                    "Unable to extract PDF page %s: %s",
                    page_number,
                    exc,
                )

                continue

            normalized_text = (
                page_text
                .replace("\x00", "")
                .strip()
            )

            if normalized_text:

                pages.append(
                    normalized_text
                )

        content = (
            "\n\n".join(pages)
            .strip()
        )

        if not content:

            raise ValueError(
                "No readable text was found in the policy PDF. "
                "This PDF may contain scanned images instead "
                "of selectable text."
            )

        return content

    # ============================================================
    # DOCX EXTRACTION
    # ============================================================

    def _extract_docx(
        self,
        document_path: Path,
    ) -> str:
        """
        Extract text from DOCX paragraphs and tables.
        """

        self._validate_file_exists(
            document_path
        )

        if (
            document_path.suffix.lower()
            != ".docx"
        ):
            raise ValueError(
                "The policy document must be a DOCX file."
            )

        try:

            document = Document(
                str(document_path)
            )

        except Exception as exc:

            logger.exception(
                "Unable to open policy DOCX: %s",
                document_path,
            )

            raise ValueError(
                "The policy DOCX could not be opened."
            ) from exc

        content_parts: list[str] = []

        # --------------------------------------------------------
        # PARAGRAPHS
        # --------------------------------------------------------

        for paragraph in document.paragraphs:

            text = (
                paragraph.text
                .replace("\x00", "")
                .strip()
            )

            if text:

                content_parts.append(
                    text
                )

        # --------------------------------------------------------
        # TABLES
        # --------------------------------------------------------

        for table in document.tables:

            for row in table.rows:

                row_values: list[str] = []

                for cell in row.cells:

                    cell_text = (
                        cell.text
                        .replace("\x00", "")
                        .strip()
                    )

                    if cell_text:

                        row_values.append(
                            cell_text
                        )

                if row_values:

                    content_parts.append(
                        " | ".join(row_values)
                    )

        content = (
            "\n\n".join(
                content_parts
            )
            .strip()
        )

        if not content:

            raise ValueError(
                "No readable text was found in the policy DOCX."
            )

        return content

    # ============================================================
    # RESOLVE DOCUMENT PATH
    # ============================================================

    def _resolve_document_path(
        self,
        document_url: str,
    ) -> Path:
        """
        Resolve a policy URL/path to a physical file.

        Security:
            Upload-relative paths are restricted to
            POLICY_UPLOAD_DIRECTORY.
        """

        normalized_url = (
            document_url or ""
        ).strip()

        if not normalized_url:

            raise ValueError(
                "Policy document URL is required."
            )

        parsed = urlparse(
            normalized_url
        )

        # ========================================================
        # HTTP / HTTPS
        # ========================================================

        if parsed.scheme in {
            "http",
            "https",
        }:

            raw_path = parsed.path

            filename = Path(
                unquote(raw_path)
            ).name

            if not filename:

                raise ValueError(
                    "Policy document filename is invalid."
                )

            upload_root = (
                Path(
                    settings.POLICY_UPLOAD_DIRECTORY
                )
                .resolve()
            )

            document_path = (
                upload_root / filename
            ).resolve()

            self._ensure_inside_directory(
                document_path,
                upload_root,
            )

            return document_path

        # ========================================================
        # FILE URL
        # ========================================================

        if (
            parsed.scheme
            == self._SUPPORTED_FILE_SCHEME
        ):

            raw_path = unquote(
                parsed.path
            )

            # Windows:
            #
            # file:///C:/folder/file.pdf
            #

            if (
                len(raw_path) >= 3
                and raw_path[0] == "/"
                and raw_path[2] == ":"
            ):

                raw_path = raw_path[1:]

            return Path(
                raw_path
            ).resolve()

        # ========================================================
        # RELATIVE PATH
        # ========================================================

        raw_path = unquote(
            normalized_url
        )

        normalized_path = (
            raw_path
            .replace("\\", "/")
            .lstrip("/")
        )

        upload_root = (
            Path(
                settings.POLICY_UPLOAD_DIRECTORY
            )
            .resolve()
        )

        # ========================================================
        # /uploads/policies/file.ext
        # ========================================================

        if normalized_path.startswith(
            "uploads/policies/"
        ):

            filename = Path(
                normalized_path
            ).name

            if not filename:

                raise ValueError(
                    "Policy document filename is invalid."
                )

            document_path = (
                upload_root / filename
            ).resolve()

            self._ensure_inside_directory(
                document_path,
                upload_root,
            )

            return document_path

        # ========================================================
        # FILENAME ONLY
        # ========================================================

        filename = Path(
            normalized_path
        ).name

        if filename:

            candidate = (
                upload_root / filename
            ).resolve()

            if candidate.exists():

                self._ensure_inside_directory(
                    candidate,
                    upload_root,
                )

                return candidate

        # ========================================================
        # DIRECT FILESYSTEM PATH
        # ========================================================

        return Path(
            normalized_path
        ).resolve()

    # ============================================================
    # PDF PATH
    # ============================================================

    def _resolve_pdf_path(
        self,
        document_url: str,
    ) -> Path:

        path = (
            self._resolve_document_path(
                document_url
            )
        )

        if (
            path.suffix.lower()
            != ".pdf"
        ):

            raise ValueError(
                "The policy document must be a PDF file."
            )

        return path

    # ============================================================
    # FILE VALIDATION
    # ============================================================

    @staticmethod
    def _validate_file_exists(
        file_path: Path,
    ) -> None:

        if not file_path.exists():

            raise FileNotFoundError(
                f"Policy document not found: {file_path}"
            )

        if not file_path.is_file():

            raise FileNotFoundError(
                f"Policy document is not a file: {file_path}"
            )

    # ============================================================
    # SECURITY
    # ============================================================

    @staticmethod
    def _ensure_inside_directory(
        file_path: Path,
        directory: Path,
    ) -> None:

        try:

            file_path.relative_to(
                directory
            )

        except ValueError as exc:

            raise ValueError(
                "Invalid policy document path."
            ) from exc